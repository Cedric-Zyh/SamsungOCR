from datetime import date, datetime
from io import BytesIO

from openpyxl import load_workbook
import pytest

from receipt_ocr.web import sample_statistics_export as exports
from test_sample_center import samples


def sheet(response):
    assert response.status_code == 200
    assert response.mimetype == 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
    assert 'attachment;' in response.headers['Content-Disposition']
    workbook = load_workbook(BytesIO(response.data))
    assert workbook.sheetnames == ['统计总览', '样单明细']
    assert workbook.active.title == '统计总览'
    return workbook['样单明细']


def details(worksheet):
    rows = list(worksheet.values)
    header = next(index for index, row in enumerate(rows) if row[0] == '样单编号')
    return rows[header], rows[header + 1:]


SECTION_TITLES = {
    '各识别方案表现', '覆盖与标注', '错误类型', '人工问题标记',
    '按公司统计 · 覆盖与效果', '按公司统计 · 错误细分',
    '按签章要求统计 · 覆盖与效果', '按签章要求统计 · 错误细分',
    '共同样本对比', '判定分歧', '日期分组',
}


def section(workbook, title):
    rows = list(workbook['统计总览'].values)
    start = next(index for index, row in enumerate(rows) if row[0] == title)
    end = next((index for index in range(start + 1, len(rows)) if rows[index][0] in SECTION_TITLES), len(rows))
    output = []
    for row in rows[start + 1:end]:
        values = list(row)
        while values and values[-1] is None:
            values.pop()
        if values:
            output.append(tuple(values))
    return output


def test_download_company_date_and_batch_scope_without_touching_production(samples):
    ctx, client = samples
    worksheet = sheet(client.get('/api/samples/statistics/export?company=甲公司&tab=date'))
    assert worksheet.title == '样单明细'
    header, rows = details(worksheet)
    assert {row[1] for row in rows} == {'one.png', 'two.png'}
    assert header[4] == '标注日期'
    assert all(row[5] == '未测试' for row in rows)
    assert worksheet.freeze_panes and worksheet.auto_filter.ref
    task = client.post('/api/samples/run', json={'filenames': ['one.png']}).json['task_id']
    scoped = sheet(client.get('/api/samples/statistics/export', query_string={'task_id': task}))
    assert [row[1] for row in details(scoped)[1]] == ['one.png']
    assert ctx.database.list_results() == [] and ctx.database.list_tasks() == []


def test_download_all_matching_rows_and_keep_untrusted_text_as_text(samples, monkeypatch):
    _, client = samples
    def c(state='正确', value='公司', decision='匹配'):
        return dict(state=state, value=value, decision=decision, note='')
    rows = [dict(filename=f'{7000000000 + index}.jpg', company='客户甲', job_status='succeeded',
                 requirement='=1+1', expected_seal='匹配', expected_date=None, date=c(),
                 seals={'qingtong_template': c('错误', '=HYPERLINK("bad")', '不匹配'),
                        'qingtong_ocr': c(), 'local:paddle_v6': c()}) for index in range(13)]
    data = dict(items=rows, columns=[{'key': 'qingtong_template', 'label': '清瞳模板'}], batches=[])
    monkeypatch.setattr(exports, 'sample_statistics', lambda *args: data)
    worksheet = sheet(client.get('/api/samples/statistics/export?filter=remote'))
    _, result = details(worksheet)
    assert len(result) == 13  # All pages, not just the visible five rows.
    assert result[0][0] == '7000000000'
    overview = worksheet.parent['统计总览']
    summary = next(row for row in overview if row[0].value == '清瞳模板' and row[1].number_format == '0.0%')
    assert summary[1].value == 0
    for row in worksheet:
        for item in row:
            assert item.data_type != 'f'
    assert result[0][4] == '=1+1'
    assert result[0][7] == '=HYPERLINK("bad")'
    assert len(details(sheet(client.get('/api/samples/statistics/export?filter=local')))[1]) == 13
    assert details(sheet(client.get('/api/samples/statistics/export?filter=missing')))[1] == []


@pytest.fixture
def evaluated_rows(monkeypatch):
    def c(state='正确', decision='匹配'):
        return dict(state=state, value='印章内容', decision=decision, note='')
    rows = []
    scenarios = [('不匹配', '错误', '匹配'), ('匹配', '错误', '不匹配'),
                 ('匹配', '错误', '部分匹配'), ('匹配', '错误', ''),
                 ('匹配', '未识别', ''), ('匹配', '失败', ''),
                 (None, '待标注', '匹配'), ('匹配', '未测试', ''), ('匹配', '正确', '匹配')]
    date_scenarios = [('正确', '2026-10-24'), ('错误', '2026-10-24'),
                      ('未识别', '2026-10-24'), ('错误', ''), ('正确', ''),
                      ('失败', '2026-10-24'), ('待标注', None), ('未测试', '2026-10-24'),
                      ('正确', '2026-10-24')]
    for index, ((expected, state, decision), (date_state, expected_date)) in enumerate(zip(scenarios, date_scenarios)):
        rows.append(dict(filename=f'{7000000000 + index}.jpg', company='客户甲' if index < 5 else '客户乙',
                         requirement='甲专用章' if index < 4 else '', expected_seal=expected,
                         expected_date=expected_date, job_status='' if index == 7 else 'failed' if index == 5 else 'succeeded',
                         tested_at='2026-10-24T00:01:00+08:00' if index != 7 else '',
                         task_id='test-batch' if index != 7 else '', task_name='样单测试 · 8 张' if index != 7 else '',
                         date=c(date_state), seals={'qingtong_template': c(state, decision),
                                                  'qingtong_ocr': c('未测试' if index == 7 else '待标注' if index == 6 else '正确', expected or '')}))
    data = dict(items=rows, columns=[{'key': 'qingtong_template', 'label': '清瞳模板'},
                                    {'key': 'qingtong_ocr', 'label': '清瞳文字'}], batches=[])
    monkeypatch.setattr(exports, 'sample_statistics', lambda *args: data)
    return rows


@pytest.mark.parametrize('selected,indices', [
    ('false_accept', [0]), ('false_reject', [1]), ('partial', [2]), ('other', [3]),
    ('wrong', [0, 1, 2, 3]), ('unrecognized', [4]), ('failed', [5]),
    ('pending_truth', [6]), ('untested', [7]), ('missing', [4, 5]),
    ('errors', [0, 1, 2, 3, 4, 5]),
])
def test_export_new_seal_filters(samples, evaluated_rows, selected, indices):
    _, client = samples
    worksheet = sheet(client.get('/api/samples/statistics/export', query_string={
        'filter': selected, 'channel': 'qingtong_template'}))
    _, rows = details(worksheet)
    assert [row[1] for row in rows] == [evaluated_rows[index]['filename'] for index in indices]


@pytest.mark.parametrize('selected,indices', [
    ('wrong_date', [1]), ('date_missed', [2]), ('date_spurious', [3]),
    ('date_with_value', [0, 1, 2, 5, 7, 8]), ('date_blank', [3, 4]),
    ('failed', [5]), ('pending_truth', [6]), ('untested', [7]),
])
def test_export_date_cohorts_and_errors(samples, evaluated_rows, selected, indices):
    _, client = samples
    worksheet = sheet(client.get('/api/samples/statistics/export', query_string={
        'tab': 'date', 'filter': selected, 'channel': 'date'}))
    assert [row[1] for row in details(worksheet)[1]] == [evaluated_rows[index]['filename'] for index in indices]


def test_export_channel_requirement_and_company_match_the_drilled_down_rows(samples, evaluated_rows):
    _, client = samples
    response = client.get('/api/samples/statistics/export', query_string={
        'company': '客户甲', 'requirement': '甲专用章', 'filter': 'errors', 'channel': 'qingtong_template'})
    worksheet = sheet(response)
    header, rows = details(worksheet)
    assert len(rows) == 4
    assert '清瞳模板 · 评估结果' in header
    assert '清瞳文字 · 评估结果' not in header
    assert rows[0][header.index('测试时间')] == datetime(2026, 10, 24, 0, 1)
    assert rows[0][header.index('测试批次 ID')] == 'test-batch'
    other = sheet(client.get('/api/samples/statistics/export', query_string={
        'company': '客户甲', 'requirement': '甲专用章', 'filter': 'errors', 'channel': 'qingtong_ocr'}))
    assert details(other)[1] == []


def test_export_coverage_annotations_error_types_and_groups(samples, evaluated_rows):
    _, client = samples
    response = client.get('/api/samples/statistics/export')
    assert response.status_code == 200
    workbook = load_workbook(BytesIO(response.data))
    assert workbook.sheetnames == ['统计总览', '样单明细']
    coverage = {row[0]: row[1] for row in section(workbook, '覆盖与标注') if row[0]}
    assert {key: coverage[key] for key in ('样单总数', '已测试', '未测试', '已标注', '待标注', '应匹配', '应不匹配')} == {
        '样单总数': 9, '已测试': 8, '未测试': 1, '已标注': 8, '待标注': 1, '应匹配': 7, '应不匹配': 1}
    assert coverage['可评测样单'] == 7 and coverage['异常样单'] == 6
    error_rows = [row for row in section(workbook, '错误类型') if row[0] == '清瞳模板']
    assert error_rows[0][:9] == ('清瞳模板', 7, 6, 1, 1, 1, 1, 1, 1)
    company_rows = [row for row in section(workbook, '按公司统计 · 覆盖与效果') if row[0] == '客户甲']
    template = next(row for row in company_rows if row[1] == '清瞳模板')
    assert template[:9] == ('客户甲', '清瞳模板', 5, 5, 0, 5, 0, 0, 5)
    assert template[-1] == 5  # Each anomalous sample counts once, even if several channels failed.
    requirement_rows = [row for row in section(workbook, '按签章要求统计 · 覆盖与效果') if row[0] == '甲专用章']
    assert all(row[2] == 4 for row in requirement_rows)
    for worksheet in workbook:
        assert all(item.data_type != 'f' for row in worksheet for item in row)


def test_date_export_breaks_down_wrong_missing_and_spurious_dates(samples, evaluated_rows):
    _, client = samples
    response = client.get('/api/samples/statistics/export?tab=date')
    workbook = load_workbook(BytesIO(response.data))
    rows = list(section(workbook, '错误类型'))
    header = next(row for row in rows if row[:2] == ('识别通道', '可评测'))
    result = next(row for row in rows if row[0] == '识别日期')
    assert header[:7] == ('识别通道', '可评测', '错误合计', '日期识别错误', '有日期但未识别', '空白日期误识别', '失败')
    assert result[:7] == ('识别日期', 7, 4, 1, 1, 1, 1)
    coverage = {row[0]: row[1] for row in section(workbook, '覆盖与标注') if row[0]}
    assert coverage['有日期'] == 6 and coverage['日期空白'] == 2
    assert workbook.sheetnames == ['统计总览', '样单明细']
    cohorts = {row[0]: row for row in section(workbook, '日期分组') if row[0] in {'有日期', '空白日期'}}
    assert cohorts['有日期'][:8] == ('有日期', 6, 5, 2, 1, 1, 1, 0.4)
    assert cohorts['空白日期'][:8] == ('空白日期', 2, 2, 1, 1, 0, 0, 0.5)


def test_export_human_problem_markers_are_separate_deduplicated_and_literal(samples, evaluated_rows):
    _, client = samples
    baseline = load_workbook(BytesIO(client.get('/api/samples/statistics/export').data))
    evaluated_rows[0].update(problem_review={'status': 'issues', 'types': ['日期读错', '日期读错', '印章判定错误']},
                             sample_features=['遮挡', '遮挡'], human_note='=1+1', legacy_problem_types=['旧印章问题'])
    evaluated_rows[1].update(problem_review={'status': 'issues', 'types': ['日期读错']}, sample_features=['遮挡'])
    evaluated_rows[2].update(problem_review={'status': 'clear', 'types': ['不应计入']})
    workbook = load_workbook(BytesIO(client.get('/api/samples/statistics/export').data))
    header, rows = details(workbook['样单明细'])
    assert header[-5:] == ('本次问题标注状态', '人工问题类型', '图片与版式特征', '补充备注', '未关联旧标记')
    assert rows[0][-5:] == ('已发现问题', '日期读错；印章判定错误', '遮挡', '=1+1', '旧印章问题')
    assert rows[2][-5:-3] == ('已核对无问题', None)
    assert rows[3][-5] == '尚未标注'
    manual = list(section(workbook, '人工问题标记'))
    assert next(row for row in manual if row[0] == '已发现问题')[1] == 2
    assert next(row for row in manual if row[0] == '日期读错')[1] == 2
    assert next(row for row in manual if row[0] == '印章判定错误')[1] == 1
    assert next(row for row in manual if row[0] == '遮挡')[1] == 2
    assert next(row for row in manual if row[0] == '有未关联旧标记')[1] == 1
    assert not any(row[0] == '不应计入' for row in manual)
    assert list(section(workbook, '错误类型')) == list(section(baseline, '错误类型'))
    assert all(item.data_type != 'f' for worksheet in workbook for row in worksheet for item in row)


def test_export_keeps_remote_local_and_unspecified_seal_markers_distinct(samples, evaluated_rows):
    _, client = samples
    baseline = load_workbook(BytesIO(client.get('/api/samples/statistics/export').data))
    evaluated_rows[0]['problem_review'] = {'status': 'issues', 'types': [
        '清瞳印章判定错误', '本地印章判定错误', '清瞳印章判定错误',
    ]}
    evaluated_rows[1]['problem_review'] = {'status': 'issues', 'types': ['印章判定错误']}
    workbook = load_workbook(BytesIO(client.get('/api/samples/statistics/export').data))
    _, rows = details(workbook['样单明细'])
    assert rows[0][-4] == '清瞳印章判定错误；本地印章判定错误'
    assert rows[1][-4] == '印章判定错误'
    manual = list(section(workbook, '人工问题标记'))
    assert next(row for row in manual if row[0] == '已发现问题')[1] == 2
    for kind in ('清瞳印章判定错误', '本地印章判定错误', '印章判定错误'):
        assert next(row for row in manual if row[0] == kind)[1] == 1
    assert list(section(workbook, '错误类型')) == list(section(baseline, '错误类型'))
    assert list(section(workbook, '共同样本对比')) == list(section(baseline, '共同样本对比'))


def test_export_coverage_separates_pending_and_trims_requirement_groups(samples, evaluated_rows):
    _, client = samples
    evaluated_rows[0]['job_status'] = 'queued'
    evaluated_rows[0]['seals'] = {'qingtong_template': {'state': '测试中'}, 'qingtong_ocr': {'state': '测试中'}}
    evaluated_rows[0]['requirement'] = ' 甲专用章 '
    response = client.get('/api/samples/statistics/export?requirement=%20甲专用章%20')
    workbook = load_workbook(BytesIO(response.data))
    assert len(details(workbook['样单明细'])[1]) == 4
    coverage = {row[0]: row[1] for row in section(workbook, '覆盖与标注') if row[0]}
    assert coverage['已测试'] == 3 and coverage['测试中'] == 1 and coverage['未测试'] == 0
    assert coverage['可评测样单'] == coverage['异常样单'] == 3
    grouped = [row for row in section(workbook, '按签章要求统计 · 覆盖与效果') if row[0] == '甲专用章']
    assert len(grouped) == 2 and all(row[2:5] == (4, 3, 0) for row in grouped)


def test_date_error_filters_treat_empty_wrong_reading_as_missed(evaluated_rows):
    missed = evaluated_rows[1]
    missed['date']['value'] = ''
    assert not exports.matches(missed, 'date', 'wrong_date')
    assert exports.matches(missed, 'date', 'date_missed')
    blank = evaluated_rows[3]
    blank['date']['value'] = ''
    assert not exports.matches(blank, 'date', 'date_spurious')


def test_common_sample_export_uses_all_active_channels_and_counts_missing_failures(samples, evaluated_rows, monkeypatch):
    _, client = samples
    remote_a, remote_b, local, dormant = 'qingtong_template', 'qingtong_ocr', 'local:model', 'local:dormant'
    scenarios = [('正确', '错误', '正确'), ('正确', '未运行', '正确'),
                 ('失败', '未识别', '未识别'), ('未运行', '正确', '正确'), ('正确', '正确', '未运行')]
    rows = evaluated_rows[:5]
    for row, states in zip(rows, scenarios):
        row['seals'] = {key: {'state': state} for key, state in zip((remote_a, remote_b, local), states)}
        row['seals'][dormant] = {'state': '未运行'}
    columns = [{'key': key, 'label': label} for key, label in (
        (remote_a, '清瞳模板'), (remote_b, '清瞳文字'), (local, '本地模型'), (dormant, '未启用模型'))]
    monkeypatch.setattr(exports, 'sample_statistics', lambda *args: dict(items=rows, columns=columns, batches=[]))
    response = client.get('/api/samples/statistics/export')
    workbook = load_workbook(BytesIO(response.data))
    common = {row[0]: row for row in section(workbook, '共同样本对比') if row[0] in {column['label'] for column in columns}}
    assert set(common) == {'清瞳模板', '清瞳文字', '本地模型'}
    assert common['清瞳模板'][:8] == ('清瞳模板', 2, 1, 0, 0, 1, 0.5, 3)
    assert common['清瞳文字'][:8] == ('清瞳文字', 2, 0, 1, 1, 0, 0, 3)
    assert common['本地模型'][:8] == ('本地模型', 2, 1, 0, 1, 0, 0.5, 3)


def test_disagreement_export_classifies_only_differing_valid_decisions(samples, evaluated_rows, monkeypatch):
    _, client = samples
    remote_a, remote_b, local = 'qingtong_template', 'qingtong_ocr', 'local:model'
    def c(state, decision):
        return {'state': state, 'decision': decision}
    scenarios = [
        [c('正确', '匹配'), c('错误', '不匹配'), c('正确', '匹配')],
        [c('错误', '匹配'), c('正确', '不匹配'), c('正确', '不匹配')],
        [c('错误', '不匹配'), c('错误', '部分匹配'), c('错误', '部分匹配')],
        [c('正确', '匹配'), c('正确', '匹配'), c('未运行', '')],
        [c('失败', '匹配'), c('错误', '不匹配'), c('未识别', '匹配')],
        [c('待标注', '匹配'), c('待标注', '不匹配'), c('待标注', '不匹配')],
    ]
    rows = evaluated_rows[:6]
    for row, cells in zip(rows, scenarios):
        row['seals'] = dict(zip((remote_a, remote_b, local), cells))
    columns = [{'key': key, 'label': label} for key, label in (
        (remote_a, '清瞳模板'), (remote_b, '清瞳文字'), (local, '本地模型'))]
    monkeypatch.setattr(exports, 'sample_statistics', lambda *args: dict(items=rows, columns=columns, batches=[]))
    workbook = load_workbook(BytesIO(client.get('/api/samples/statistics/export').data))
    pairs = {(row[0], row[1]): row[2:8] for row in section(workbook, '判定分歧') if row[0] in {'清瞳模板', '清瞳文字'}}
    assert pairs[('清瞳模板', '清瞳文字')] == (4, 3, 1, 1, 1, 0)
    assert pairs[('清瞳模板', '本地模型')] == (3, 2, 0, 1, 1, 0)
    assert pairs[('清瞳文字', '本地模型')] == (3, 1, 0, 1, 0, 0)


def test_group_export_prioritizes_anomalous_rows_over_group_size(samples, evaluated_rows):
    _, client = samples
    for row in evaluated_rows:
        row['company'] = '样单较多'
        row['seals'] = {'qingtong_template': {'state': '正确'}, 'qingtong_ocr': {'state': '正确'}}
    evaluated_rows[0]['company'] = '异常集中'
    evaluated_rows[0]['seals'] = {'qingtong_template': {'state': '错误'}, 'qingtong_ocr': {'state': '失败'}}
    workbook = load_workbook(BytesIO(client.get('/api/samples/statistics/export').data))
    groups = [row for row in section(workbook, '按公司统计 · 覆盖与效果') if row[0] in {'异常集中', '样单较多'}]
    assert groups[0][0] == '异常集中' and groups[0][-1] == 1
    assert groups[-1][0] == '样单较多' and groups[-1][-1] == 0


def test_redesigned_report_keeps_overview_compact_and_detail_filterable(samples, evaluated_rows):
    _, client = samples
    workbook = load_workbook(BytesIO(client.get('/api/samples/statistics/export').data))
    assert workbook.sheetnames == ['统计总览', '样单明细']
    assert workbook.active.title == '统计总览'
    overview, register = workbook['统计总览'], workbook['样单明细']
    assert overview.max_column == 12
    assert not overview.sheet_view.showGridLines and not register.sheet_view.showGridLines
    assert isinstance(overview['B5'].value, datetime)
    assert overview.freeze_panes and register.freeze_panes == 'D9'
    assert register.auto_filter.ref.startswith('A8:')
    assert register.print_title_rows == '$1:$8'
    assert overview.print_area and register.print_area
    assert overview['A8'].value == 9 and overview['B8'].value == 8
    assert overview['E8'].value == 7 and overview['I8'].value == 6
    assert overview['B12'].number_format == '0.0%'
    assert register['A9'].data_type == 's' and register['A9'].number_format == '@'
    assert register['G9'].fill.fgColor.rgb == '00F8E9E1'
    assert register.row_dimensions[9].height <= 60
    granular = section(workbook, '按公司统计 · 错误细分')
    header = granular[0]
    template = next(row for row in granular if row[:2] == ('客户甲', '清瞳模板'))
    assert {name: template[header.index(name)] for name in ('误放行', '误拒绝', '部分匹配', '其他错误', '未识别', '失败')} == {
        '误放行': 1, '误拒绝': 1, '部分匹配': 1, '其他错误': 1, '未识别': 1, '失败': 0}


def test_dates_are_excel_dates_and_timestamps_use_shanghai_time(samples, evaluated_rows):
    _, client = samples
    evaluated_rows[0]['date']['value'] = '2026-10-24'
    evaluated_rows[0]['tested_at'] = '2026-10-23T16:01:00Z'
    register = sheet(client.get('/api/samples/statistics/export?tab=date'))
    header, rows = details(register)
    for name in ('标注日期', '识别日期 · 识别内容'):
        index = header.index(name)
        assert rows[0][index] == datetime(2026, 10, 24)
        assert register.cell(9, index + 1).number_format == 'yyyy-mm-dd'
    index = header.index('测试时间')
    assert rows[0][index] == datetime(2026, 10, 24, 0, 1)
    assert register.cell(9, index + 1).number_format == 'yyyy-mm-dd hh:mm:ss'
    assert rows[3][header.index('标注日期')] == '空白'
    assert rows[6][header.index('标注日期')] == '待标注'


@pytest.mark.parametrize('query,status', [
    ('task_id=unknown', 404), ('tab=unknown', 400), ('filter=unknown', 400),
    ('tab=date&filter=local', 400),
    ('channel=unknown', 400), ('tab=date&channel=qingtong_template', 400),
    ('tab=seal&filter=wrong_date', 400), ('tab=date&filter=false_accept', 400),
])
def test_invalid_export_scope_reports_error(samples, query, status):
    _, client = samples
    response = client.get('/api/samples/statistics/export?' + query)
    assert response.status_code == status
    assert response.json['error']
