"""Excel download for the sample statistics view, including every filtered row."""
from io import BytesIO
import re

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from .sample_statistics import sample_statistics


BAD = {'错误', '未识别', '失败'}
SCORED = BAD | {'正确'}
FILTERS = {'': '全部', 'errors': '错误', 'remote': '清瞳两路不一致',
           'local': '清瞳与本地不一致', 'missing': '未识别 / 失败',
           'false_accept': '误放行', 'false_reject': '误拒绝', 'partial': '部分匹配',
           'pending_truth': '待标注', 'untested': '未测试', 'failed': '失败',
           'wrong': '判定错误', 'unrecognized': '未识别', 'other': '其他错误',
           'wrong_date': '日期识别错误', 'date_missed': '有日期但未识别',
           'date_spurious': '空白日期误识别', 'date_with_value': '已标注有日期',
           'date_blank': '已标注日期空白'}
NOTE = '未识别、失败计入正确率分母；未测试、未运行、测试中、待标注不计入。印章按人工标注的应匹配 / 应不匹配评估判定，不代表逐字转写准确率。'
STATES = ('正确', '错误', '未识别', '失败', '待标注', '未运行', '未测试', '测试中')
SEAL_FILTERS = {'remote', 'local', 'false_accept', 'false_reject', 'partial', 'other'}
DATE_FILTERS = {'wrong_date', 'date_missed', 'date_spurious', 'date_with_value', 'date_blank'}


def cell(row, key):
    return row['date'] if key == 'date' else row['seals'].get(key, {
        'state': '未运行' if row['job_status'] else '未测试'})


def disagrees(cells):
    return len({c.get('decision') for c in cells if c['state'] in SCORED
                and c['state'] not in {'未识别', '失败'}
                and c.get('decision') in {'匹配', '不匹配', '部分匹配'}}) > 1


def error_type(row, current):
    if current['state'] != '错误':
        return ''
    expected, decision = row.get('expected_seal'), current.get('decision')
    if expected == '不匹配' and decision == '匹配':
        return 'false_accept'
    if expected == '匹配' and decision == '不匹配':
        return 'false_reject'
    return 'partial' if decision == '部分匹配' else 'other'


def matches(row, tab, selected, channel=''):
    cells = [cell(row, channel)] if channel else (
        [row['date']] if tab == 'date' else list(row['seals'].values()))
    if selected == 'pending_truth':
        return row['expected_date' if tab == 'date' else 'expected_seal'] is None
    if selected == 'untested':
        return not row['job_status']
    if selected == 'date_with_value':
        return bool(row['expected_date'])
    if selected == 'date_blank':
        return row['expected_date'] == ''
    if selected in {'false_accept', 'false_reject', 'partial', 'other'}:
        return any(error_type(row, current) == selected for current in cells)
    if selected in {'wrong_date', 'date_missed', 'date_spurious'}:
        current = row['date']
        if selected == 'date_spurious':
            return current['state'] == '错误' and row['expected_date'] == '' and bool(current.get('value'))
        if selected == 'date_missed':
            return bool(row['expected_date']) and (current['state'] == '未识别'
                or (current['state'] == '错误' and not current.get('value')))
        return bool(row['expected_date']) and current['state'] == '错误' and bool(current.get('value'))
    if selected == 'errors':
        return any(c['state'] in BAD for c in cells)
    if selected == 'failed':
        return any(c['state'] == '失败' for c in cells)
    if selected in {'wrong', 'unrecognized'}:
        return any(c['state'] == ('错误' if selected == 'wrong' else '未识别') for c in cells)
    if selected == 'missing':
        return any(c['state'] in {'未识别', '失败'} for c in cells)
    if selected == 'remote':
        return disagrees([cell(row, 'qingtong_template'), cell(row, 'qingtong_ocr')])
    if selected == 'local':
        return any(disagrees([cell(row, key), c])
                   for key in ('qingtong_template', 'qingtong_ocr')
                   for method, c in row['seals'].items() if method.startswith('local:'))
    return True


def counts(rows, key):
    states = {state: sum(cell(row, key)['state'] == state for row in rows) for state in STATES}
    states['可评测'] = sum(states[state] for state in SCORED)
    states['正确率'] = states['正确'] / states['可评测'] if states['可评测'] else None
    for kind in ('false_accept', 'false_reject', 'partial', 'other'):
        states[kind] = sum(error_type(row, cell(row, key)) == kind for row in rows) if key != 'date' else 0
    if key == 'date':
        states['other'] = states['错误']
        for kind in ('wrong_date', 'date_missed', 'date_spurious'):
            states[kind] = sum(matches(row, 'date', kind) for row in rows)
    return states


def _append(sheet, values):
    # Keep names, requirements and OCR text literal in every worksheet.
    sheet.append(values)
    for item in sheet[sheet.max_row]:
        if isinstance(item.value, str):
            item.value = re.sub(r'[\x00-\x08\x0b\x0c\x0e-\x1f]', '', item.value)
            item.data_type = 's'
        item.alignment = Alignment(vertical='top', wrap_text=True)


def _style(sheet, headers):
    for index in headers:
        for item in sheet[index]:
            item.fill = PatternFill('solid', fgColor='343D2E')
            item.font = Font(color='FFFFFF', bold=True)
        sheet.row_dimensions[index].height = 32
    sheet['A1'].font = Font(size=18, bold=True, color='343D2E')
    for index in range(1, sheet.max_column + 1):
        sheet.column_dimensions[get_column_letter(index)].width = 24 if index < 3 else 34


def _summary_sheet(workbook, title, context, header):
    sheet = workbook.create_sheet(title)
    _append(sheet, [title])
    for values in context:
        _append(sheet, values)
    _append(sheet, [])
    _append(sheet, header)
    return sheet, sheet.max_row


def _seal_comparisons(workbook, rows, columns, context):
    active = [column for column in columns if counts(rows, column['key'])['可评测']]
    common, common_header = _summary_sheet(workbook, '共同样本对比', context, [
        '识别通道', '共同可评测样单', '正确', '错误', '未识别', '失败', '正确率', '排除样单'])
    shared = [row for row in rows if all(cell(row, column['key'])['state'] in SCORED for column in active)] if len(active) >= 2 else []
    for column in active if len(active) >= 2 else []:
        summary = counts(shared, column['key'])
        _append(common, [column['label'], len(shared), *[summary[state] for state in ('正确', '错误', '未识别', '失败')],
                         summary['正确率'], len(rows) - len(shared)])
        common.cell(common.max_row, 7).number_format = '0.0%'
    _append(common, ['对比口径', '仅纳入有可评测结果的通道，并使用这些通道均可评测的同一组样单；未识别与失败计入分母。'])
    _append(common, ['时间说明', '各通道结果或配置可能来自不同测试时间；可在明细中查看测试批次。'])
    if len(active) < 2:
        _append(common, ['对比状态', '至少两个通道有可评测结果时才能进行共同样本对比。'])
    _style(common, (common_header,))
    common.freeze_panes = f'B{common_header + 1}'

    comparisons, comparison_header = _summary_sheet(workbook, '判定分歧', context, [
        '方案 A', '方案 B', '可对比样单', '判定不同', '仅 A 正确', '仅 B 正确', '双方都错', '双方均正确'])
    remote = ['qingtong_template', 'qingtong_ocr']
    pairs = [tuple(remote), *[(remote_key, column['key']) for column in columns
                             if column['key'].startswith('local:') for remote_key in remote]]
    labels = {column['key']: column['label'] for column in columns}
    for left_key, right_key in pairs:
        total = different = only_left = only_right = both_wrong = both_correct = 0
        for row in rows:
            left, right = cell(row, left_key), cell(row, right_key)
            if not all(current['state'] in {'正确', '错误'} and current.get('decision') in {'匹配', '不匹配', '部分匹配'}
                       for current in (left, right)):
                continue
            total += 1
            if left['decision'] == right['decision']:
                continue
            different += 1
            if left['state'] == right['state'] == '正确':
                both_correct += 1
            elif left['state'] == '正确':
                only_left += 1
            elif right['state'] == '正确':
                only_right += 1
            else:
                both_wrong += 1
        if total:
            _append(comparisons, [labels.get(left_key, left_key), labels.get(right_key, right_key),
                                  total, different, only_left, only_right, both_wrong, both_correct])
    _append(comparisons, ['对比口径', '双方均为正确或错误，且判定为匹配、不匹配或部分匹配时才可对比；未识别和失败不参与。'])
    _append(comparisons, ['分类口径', '仅 A 正确、仅 B 正确、双方都错、双方均正确仅统计判定不同的样单。'])
    _style(comparisons, (comparison_header,))
    comparisons.freeze_panes = f'C{comparison_header + 1}'


def _date_cohorts(workbook, rows, context):
    cohorts, header = _summary_sheet(workbook, '日期分组', context, [
        '人工标注', '样单数', '可评测', '正确', '错误', '未识别', '失败', '正确率'])
    for label, blank in (('有日期', False), ('空白日期', True)):
        members = [row for row in rows if row['expected_date'] is not None and bool(row['expected_date']) != blank]
        summary = counts(members, 'date')
        _append(cohorts, [label, len(members), summary['可评测'], *[summary[state] for state in ('正确', '错误', '未识别', '失败')],
                          summary['正确率']])
        cohorts.cell(cohorts.max_row, 8).number_format = '0.0%'
    _append(cohorts, ['评估口径', '待标注样单不纳入两类分组；未测试、未运行、测试中不计入正确率分母，未识别与失败计入。'])
    _style(cohorts, (header,))
    cohorts.freeze_panes = f'B{header + 1}'


def export_statistics(center, *, company='', requirement='', task_id='', tab='seal', selected='', channel=''):
    if (tab not in {'seal', 'date'} or selected not in FILTERS
            or (tab == 'date' and selected in SEAL_FILTERS)
            or (tab == 'seal' and selected in DATE_FILTERS)):
        raise ValueError('统计筛选条件无效')
    data = sample_statistics(center, task_id)
    columns = data['columns'] if tab == 'seal' else [{'key': 'date', 'label': '识别日期'}]
    if channel and channel not in {column['key'] for column in columns}:
        raise ValueError('统计识别通道无效')
    rows = [row for row in data['items'] if (not company or row['company'] == company)
            and (not requirement or str(row.get('requirement') or '').strip() == requirement.strip())
            and matches(row, tab, selected, channel)]
    if channel:
        columns = [column for column in columns if column['key'] == channel]
    error_metrics = [('false_accept', '误放行'), ('false_reject', '误拒绝'),
                     ('partial', '部分匹配'), ('other', '其他错误')] if tab == 'seal' else [
                     ('wrong_date', '日期识别错误'), ('date_missed', '有日期但未识别'),
                     ('date_spurious', '空白日期误识别')]
    batch = next((batch for batch in data['batches'] if batch['id'] == task_id), None)
    scope = f"{batch['name']} · {batch['created_at']} · {batch['total']} 张" if batch else '每张样单最近一次测试'
    context = [['客户公司', company or '全部公司', '统计范围', scope],
               ['签章要求', requirement or '全部签章要求', '筛选条件', FILTERS[selected]],
               ['识别通道', columns[0]['label'] if channel else '全部通道', '导出样单数', len(rows)]]
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = '印章结果' if tab == 'seal' else '日期结果'

    def append(values):
        _append(sheet, values)

    append(['样单统计结果 · ' + sheet.title])
    append(['客户公司', company or '全部公司', '统计范围', scope])
    append(['筛选条件', FILTERS[selected], '导出样单数', len(rows), '签章要求', requirement or '全部签章要求'])
    append([NOTE])
    sheet.merge_cells(start_row=4, start_column=1, end_row=4, end_column=8)
    sheet.row_dimensions[4].height = 42
    append([])
    append(['识别通道', '正确率', '正确', '可评测', '错误', '未识别', '失败', '待标注', '未运行', '未测试', '测试中',
            *[label for _, label in error_metrics]])
    for column in columns:
        summary = counts(rows, column['key'])
        append([column['label'], summary['正确率'], summary['正确'], summary['可评测'],
                *[summary[state] for state in ('错误', '未识别', '失败', '待标注', '未运行', '未测试', '测试中')],
                *[summary[key] for key, _ in error_metrics]])
        sheet.cell(sheet.max_row, 2).number_format = '0.0%'
    append([])
    header = ['样单编号', '文件名', '客户公司', '测试状态',
              *(['签章要求', '标注结论'] if tab == 'seal' else ['标注日期'])]
    for column in columns:
        header.extend(f"{column['label']} · {label}" for label in ('评估结果', '识别内容', '判定', '说明'))
    header.extend(['测试时间', '测试批次', '测试批次 ID'])
    append(header)
    header_row = sheet.max_row
    states = {'succeeded': '已完成', 'failed': '运行失败', 'queued': '测试中',
              'running': '测试中', 'awaiting_upload': '测试中'}
    from .sample_library import sample_number
    for row in rows:
        values = [sample_number(row['filename']), row['filename'], row['company'],
                  states.get(row['job_status'], '未测试')]
        values.extend([row['requirement'], f"应{row['expected_seal']}" if row['expected_seal'] else '待标注']
                      if tab == 'seal' else ['待标注' if row['expected_date'] is None else row['expected_date'] or '空白'])
        for column in columns:
            c = cell(row, column['key'])
            values.extend(c.get(key, '') for key in ('state', 'value', 'decision', 'note'))
        values.extend(row.get(key, '') for key in ('tested_at', 'task_name', 'task_id'))
        append(values)
    _style(sheet, (6, header_row))
    sheet.freeze_panes = f'D{header_row + 1}'
    sheet.auto_filter.ref = f'A{header_row}:{get_column_letter(len(header))}{sheet.max_row}'
    coverage, coverage_header = _summary_sheet(workbook, '覆盖与标注', context, ['指标', '样单数', '说明'])
    labeled_key = 'expected_date' if tab == 'date' else 'expected_seal'
    labeled = sum(row[labeled_key] is not None for row in rows)
    tested = sum(row['job_status'] in {'succeeded', 'failed'} for row in rows)
    untested = sum(not row['job_status'] for row in rows)
    evaluated = sum(any(cell(row, column['key'])['state'] in SCORED for column in columns) for row in rows)
    abnormal = sum(any(cell(row, column['key'])['state'] in BAD for column in columns) for row in rows)
    coverage_rows = [
        ['样单总数', len(rows), '按当前公司、签章要求、批次和筛选条件统计'],
        ['已测试', tested, '测试已成功完成或失败，排队和执行中另列'],
        ['未测试', untested, '没有测试任务'],
        ['已完成', sum(row['job_status'] == 'succeeded' for row in rows), '测试成功完成'],
        ['测试失败', sum(row['job_status'] == 'failed' for row in rows), '整次测试失败'],
        ['测试中', sum(row['job_status'] in {'queued', 'running', 'awaiting_upload'} for row in rows), '排队或执行中'],
        ['可评测样单', evaluated, '当前导出通道中至少一路可评测，同一样单仅计一次'],
        ['异常样单', abnormal, '当前导出通道中至少一路错误、未识别或失败，同一样单仅计一次'],
        ['已标注', labeled, '本统计对象已有人工标注'],
        ['待标注', len(rows) - labeled, '缺少人工标注，不计入正确率'],
    ]
    annotation_values = [('有日期', lambda row: bool(row['expected_date'])),
                         ('日期空白', lambda row: row['expected_date'] == '')] if tab == 'date' else [
                         ('应匹配', lambda row: row['expected_seal'] == '匹配'),
                         ('应不匹配', lambda row: row['expected_seal'] == '不匹配')]
    coverage_rows.extend([label, sum(predicate(row) for row in rows), '人工标注分布']
                         for label, predicate in annotation_values)
    for values in coverage_rows:
        _append(coverage, values)
    _append(coverage, ['评估口径', NOTE])
    if batch:
        _append(coverage, ['批次 ID', batch['id']])
        _append(coverage, ['批次状态', batch.get('status', ''),
                           f"已完成 {batch.get('completed', 0)} / {batch['total']}"])
        for stage, label in (('date', '日期方法'), ('seal', '印章方法')):
            _append(coverage, [label, '、'.join(batch.get('scheme_labels', {}).get(stage, [])) or '未运行'])
    _style(coverage, (coverage_header,))
    coverage.freeze_panes = f'B{coverage_header + 1}'

    error_tail = [('未识别', '未识别'), ('失败', '失败')] if tab == 'seal' else [('失败', '失败')]
    error_header = ['识别通道', '可评测', '错误合计', *[label for _, label in error_metrics + error_tail]]
    errors, errors_header = _summary_sheet(workbook, '错误类型', context, error_header)
    for column in columns:
        summary = counts(rows, column['key'])
        _append(errors, [column['label'], summary['可评测'], sum(summary[state] for state in BAD),
                          *[summary[key] for key, _ in error_metrics + error_tail]])
    definitions = [['误放行', '标注应不匹配，通道判定匹配'], ['误拒绝', '标注应匹配，通道判定不匹配'],
                   ['部分匹配', '已有标注，但通道仅判定部分匹配'],
                   ['其他错误', '不属于上述三类的判定错误']] if tab == 'seal' else [
                   ['日期识别错误', '已标注有日期，识别日期与标注不同'],
                   ['有日期但未识别', '已标注有日期，但通道没有识别出日期'],
                   ['空白日期误识别', '已标注日期空白，但通道识别出日期']]
    for values in definitions + [['计数口径', '各通道独立计数，同一样单在不同通道可重复出现；错误合计含未识别和失败']]:
        _append(errors, values)
    _style(errors, (errors_header,))

    for field, title, label in (('company', '按公司统计', '客户公司'), ('requirement', '按签章要求统计', '签章要求')):
        group_columns = [
            label, '样单总数', '已测试', '未测试', '已标注', '待标注', '识别通道', '正确率', '可评测',
            '正确', '错误', *[name for _, name in error_metrics], '未识别', '失败', '未运行', '测试中', '异常样单']
        grouped, group_header = _summary_sheet(workbook, title, context, group_columns)
        groups = {}
        for row in rows:
            groups.setdefault(str(row.get(field) or '').strip(), []).append(row)
        anomalies = {group: sum(any(cell(row, column['key'])['state'] in BAD for column in columns) for row in members)
                     for group, members in groups.items()}
        for group, members in sorted(groups.items(), key=lambda pair: (-anomalies[pair[0]], -len(pair[1]), pair[0])):
            group_tested = sum(row['job_status'] in {'succeeded', 'failed'} for row in members)
            group_untested = sum(not row['job_status'] for row in members)
            group_labeled = sum(row[labeled_key] is not None for row in members)
            for column in columns:
                summary = counts(members, column['key'])
                _append(grouped, [group or '未填写', len(members), group_tested, group_untested,
                                   group_labeled, len(members) - group_labeled, column['label'],
                                   summary['正确率'], summary['可评测'], summary['正确'], summary['错误'],
                                   *[summary[key] for key, _ in error_metrics],
                                   *[summary[state] for state in ('未识别', '失败', '未运行', '测试中')], anomalies[group]])
                grouped.cell(grouped.max_row, 8).number_format = '0.0%'
        _style(grouped, (group_header,))
        grouped.freeze_panes = f'B{group_header + 1}'
        grouped.auto_filter.ref = f'A{group_header}:{get_column_letter(len(group_columns))}{grouped.max_row}'
    if tab == 'seal':
        _seal_comparisons(workbook, rows, columns, context)
    else:
        _date_cohorts(workbook, rows, context)
    output = BytesIO()
    workbook.save(output)
    output.seek(0)
    return output
